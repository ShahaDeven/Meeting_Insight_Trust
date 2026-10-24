"""Gold "project state after each meeting" tables for the memory experiment.

A stronger model than the extractor (Sonnet vs Haiku) drafts each series' table from the four
human transcripts and AMI summaries; a human then verifies it in data/gold/state_tables.xlsx.
The README reports this as model-drafted, human-verified gold.

  python -m evals.state_tables                    # draft the primary series, write the workbook
  python -m evals.state_tables --series ES2004

Out: data/gold/{series}.state.draft.json and data/gold/state_tables.xlsx
"""

import argparse
import time
from typing import Literal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import BaseModel

from core.config import DATA, PRIMARY_SERIES
from core.llm import Budget, client
from core.transcript import load_transcript
from extract.extractor import render
from extract.schema import InsightType
from pipeline.ami import load_summary

GOLD = DATA / "gold"
WORKBOOK = GOLD / "state_tables.xlsx"
MODEL = "claude-sonnet-5-5"
Meeting = Literal["a", "b", "c", "d"]

SYSTEM = """You build the gold record of a project's state across a series of four meetings, \
used to test whether automatic meeting memory keeps facts right as they change.

The four-person team designs a TV remote control (PM project manager, ME marketing expert, \
UI user interface designer, ID industrial designer). You get accurate human transcripts and \
human-written summaries of meetings a-d.

List the facts a project record should hold, one row per fact slot:
- decision slots: one per design dimension the team settles (e.g. power source, case \
material, screen, buttons, target group, special features). Value = what was decided.
- cost_figure slots: one per figure (e.g. selling price, production cost limit, profit target).
- action_item slots: one per task assigned to someone. Value = "OWNER: task (open)", then \
"done" once reported back.
- open_risk slots: one per unresolved concern that matters for the design. Value = the \
concern, then "resolved: how" if it is settled later.

For each slot give its value as of the END of each meeting: null until first raised, carried \
forward unchanged, updated when it changes, "dropped" if abandoned. Only decisions the group \
actually agreed count; proposals are not decisions. Cite one short exact quote per change, \
with the meeting letter. Aim for the 20-40 slots that matter most; skip small talk and \
one-off remarks."""


class _Evidence(BaseModel):
    meeting: Meeting
    quote: str


class _Slot(BaseModel):
    slot: str
    category: InsightType
    after_a: str | None
    after_b: str | None
    after_c: str | None
    after_d: str | None
    evidence: list[_Evidence]


class StateTable(BaseModel):
    slots: list[_Slot]


def draft(series: str, budget: Budget) -> StateTable:
    parts = []
    for part in "abcd":
        mid = f"{series}{part}"
        summary = "\n".join(f"{k}: {' '.join(v) or 'NA'}" for k, v in load_summary(mid).items())
        parts.append(f"## Meeting {part}\nSummary:\n{summary}\n\nTranscript:\n"
                     f"{render(load_transcript(mid, 'clean'))}")
    budget.check()
    t0 = time.time()
    with client.messages.stream(
            model=MODEL, max_tokens=64000, system=SYSTEM,
            messages=[{"role": "user", "content": "\n\n".join(parts)}],
            output_format=StateTable, output_config={"effort": "high"},
            # Server-side fallback if a safety classifier declines (Claude API only).
            extra_headers={"anthropic-beta": "server-side-fallback-2026-07-01"},
            extra_body={"fallbacks": "default"}) as stream:
        msg = stream.get_final_message()
    budget.charge(MODEL, msg.usage, f"state:{series}", time.time() - t0)
    if msg.stop_reason != "end_turn":
        raise RuntimeError(f"{series}: stop_reason={msg.stop_reason}")
    table = getattr(msg, "parsed_output", None) or StateTable.model_validate_json(
        next(b.text for b in msg.content if b.type == "text"))
    path = GOLD / f"{series}.state.draft.json"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(table.model_dump_json(indent=1), encoding="utf-8")
    return table


def write_workbook(tables: dict[str, StateTable]) -> None:
    wb = Workbook()
    guide = wb.active
    guide.title = "How to check"
    guide.column_dimensions["A"].width = 120
    for i, line in enumerate([
        "Each series sheet is a model-drafted record of the project's state after meetings a-d.",
        "For each row: is the slot real and are the values right at the end of each meeting?",
        "  - Fix wrong values in place, then set CHECK to 'fixed'.",
        "  - If a row is right, set CHECK to 'ok'. If it should not exist, set 'delete'.",
        "  - Add missing slots in empty rows at the bottom with CHECK = 'added'.",
        "Blank = not raised yet. 'dropped' = abandoned. Values carry forward until they change.",
        "Only agreed decisions count; proposals are not decisions.",
        "Use the evidence column and the AMI summaries; open a transcript only when unsure.",
    ], 1):
        guide.cell(i, 1, line).alignment = Alignment(wrap_text=True)

    headers = ["slot", "category", "after a", "after b", "after c", "after d", "evidence",
               "CHECK", "note"]
    widths = [24, 12, 28, 28, 28, 28, 60, 10, 30]
    for series, table in tables.items():
        ws = wb.create_sheet(series)
        for c, (h, w) in enumerate(zip(headers, widths), 1):
            cell = ws.cell(1, c, h)
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="DDDDDD")
            ws.column_dimensions[cell.column_letter].width = w
        ws.freeze_panes = "C2"
        order = {t: i for i, t in enumerate(InsightType)}
        slots = sorted(table.slots, key=lambda s: order[s.category])
        for r, s in enumerate(slots, 2):
            evidence = "\n".join(f"{e.meeting}: \"{e.quote}\"" for e in s.evidence)
            for c, v in enumerate([s.slot, str(s.category), s.after_a, s.after_b, s.after_c,
                                   s.after_d, evidence], 1):
                ws.cell(r, c, v).alignment = Alignment(wrap_text=True, vertical="top")
            ws.cell(r, 8).fill = PatternFill("solid", fgColor="FFF2CC")
        check = DataValidation(type="list", formula1='"ok,fixed,delete,added"', allow_blank=True)
        ws.add_data_validation(check)
        check.add(f"H2:H{len(slots) + 30}")  # room for added rows
    wb.save(WORKBOOK)


def main() -> None:
    p = argparse.ArgumentParser()
    p.add_argument("--series", nargs="+", default=PRIMARY_SERIES)
    args = p.parse_args()
    if WORKBOOK.exists():
        raise SystemExit(f"{WORKBOOK} exists; it may hold your checks. Move it aside first.")
    budget = Budget()
    tables = {}
    for s in args.series:
        cached = GOLD / f"{s}.state.draft.json"
        tables[s] = (StateTable.model_validate_json(cached.read_text(encoding="utf-8"))
                     if cached.exists() else draft(s, budget))
        print(f"{s}: {len(tables[s].slots)} slots (total ${budget.spent:.3f})", flush=True)
    write_workbook(tables)
    print(f"Wrote {WORKBOOK}. Spent ${budget.spent:.3f}.")


if __name__ == "__main__":
    main()

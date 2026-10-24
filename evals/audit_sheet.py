"""Human audit of the judge: blind sheets, adjudication, and agreement scores.

Flow (the first audit showed the rubric was ambiguous, see CLAUDE.md):
  1. audit_sheet.xlsx      blind 100-row audit of judge v1 (done; archived in judge_v1/)
  2. --adjudicate          adjudication.xlsx: rows where human and judge v1 disagreed, with the
                           judge's reason shown; the human gives a final mark under the new rules
  3. --finalize            dev_labels.json = audit marks + adjudicated marks (the judge's dev set)
  4. --test                test_sheet.xlsx: 50 fresh blind rows, never used for tuning
  5. --score dev|test      agreement of the CURRENT judge labels with the human labels

Blind sheets never show the condition, the noisy quote or any judge verdict.
"""

import argparse
import json
import random

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from core.config import INSIGHTS
from core.jsonl import read_jsonl
from core.transcript import load_transcript
from evals.judge import LABELS, RULES, summary_text, truth_excerpt
from extract.schema import Insight

V1 = LABELS / "judge_v1"
AUDIT, AUDIT_KEY = V1 / "audit_sheet.xlsx", V1 / "audit_key.json"
ADJ, ADJ_KEY = LABELS / "adjudication.xlsx", LABELS / "adjudication_key.json"
DEV = LABELS / "dev_labels.json"
TEST, TEST_KEY = LABELS / "test_sheet.xlsx", LABELS / "test_key.json"
TEST_PER_CONDITION = {"farfield": 18, "headset": 12, "speaker_swap": 10, "clean": 10}

EXAMPLES = [
    ("action_item | owner: team", "Check the financial feasibility of adding voice recognition.",
     "UI: We have to check into the financial feasibility whether we can incorporate it...\n"
     "ME: ...there was just a cost issue with that, we just need to check on the cost.",
     "wrong", "owner", "UI committed to checking it, not the whole team."),
    ("decision", "Start with a black and white LCD screen.",
     "UI: Maybe we could start with the black and white.\nID: Yeah.",
     "wrong", "type", "A suggestion answered only by 'yeah' is a proposal, not a decision."),
    ("decision", "The remote will control TV and video (universal).",
     "PM: we decided upon a universal control, one handset for all...\n"
     "(summary: the team later restricted it to a TV-only remote)",
     "wrong", "content", "Said, but superseded later in the meeting: a stale fact."),
]

def _load_insights() -> dict[str, Insight]:
    out = {}
    for f in INSIGHTS.glob("*.r*.jsonl"):
        if not f.name.endswith(".rejected.jsonl"):
            out.update({i.id: i for i in read_jsonl(f, Insight)})
    return out


def current_verdicts() -> dict[str, dict]:
    """Latest judge labels (data/labels/*.judge.jsonl), keyed by insight id."""
    return {r["insight_id"]: r for f in LABELS.glob("*.judge.jsonl") for r in read_jsonl(f)}


def _guide(wb: Workbook, intro: str) -> None:
    ws = wb.active
    ws.title = "How to mark"
    ws.column_dimensions["A"].width = 120
    lines = [intro, "", *RULES.splitlines(), "",
             "Mark: correct / wrong / unsure. If wrong, pick why: content, type, owner, number.",
             "The meeting summary column is the human-written AMI summary: use it to spot facts "
             "that were reversed or superseded later in the meeting.", "", "Worked examples:"]
    for kind, claim, said, mark, why, note in EXAMPLES:
        lines += [f"  {kind}: {claim}", *(f"    {s}" for s in said.splitlines()),
                  f"    -> {mark} / {why}: {note}", ""]
    for i, line in enumerate(lines, 1):
        ws.cell(i, 1, line).alignment = Alignment(wrap_text=True)
    ws.cell(1, 1).font = Font(bold=True)


def _sheet(wb: Workbook, title: str, headers: list[str], widths: list[int]):
    ws = wb.create_sheet(title)
    for c, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(1, c, h)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDDDDD")
        ws.column_dimensions[cell.column_letter].width = w
    ws.freeze_panes = "A2"
    return ws


def _row_values(ins: Insight, cache: dict) -> list:
    clean = cache.setdefault(ins.meeting_id, load_transcript(ins.meeting_id, "clean"))
    details = "; ".join(f"{k}: {v}" for k, v in
                        [("owner", ins.owner or ("none named" if ins.type == "action_item" else None)),
                         ("value", ins.value), ("deadline", ins.deadline)] if v)
    return [str(ins.type), details, ins.summary, truth_excerpt(clean, ins), summary_text(ins.meeting_id)]


def _write_rows(ws, rows: list[list], mark_col: int) -> None:
    for r, values in enumerate(rows, 2):
        for c, v in enumerate(values, 1):
            ws.cell(r, c, v).alignment = Alignment(wrap_text=True, vertical="top")
        for c in (mark_col, mark_col + 1):
            ws.cell(r, c).fill = PatternFill("solid", fgColor="FFF2CC")
    for col, options in ((mark_col, "correct,wrong,unsure"), (mark_col + 1, "content,type,owner,number")):
        dv = DataValidation(type="list", formula1=f'"{options}"', allow_blank=True)
        ws.add_data_validation(dv)
        letter = ws.cell(1, col).column_letter
        dv.add(f"{letter}2:{letter}{len(rows) + 1}")


def read_marks(path, n: int, mark_col: int) -> dict[int, tuple[str, str, str]]:
    ws = load_workbook(path, read_only=True).worksheets[1]
    marks = {}
    # Bound the scan: Excel can save formatting down to row 1,048,576.
    for row in ws.iter_rows(min_row=2, max_row=n + 1, max_col=mark_col + 2, values_only=True):
        clean = lambda v: (v or "").strip().lower() if isinstance(v, str) else ""  # noqa: E731
        marks[row[0]] = (clean(row[mark_col - 1]), clean(row[mark_col]), row[mark_col + 1] or "")
    return marks


def adjudicate() -> None:
    if ADJ.exists():
        raise SystemExit(f"{ADJ} exists; delete it first for a fresh one.")
    key = json.loads(AUDIT_KEY.read_text(encoding="utf-8"))
    marks = read_marks(AUDIT, len(key), 6)
    v1 = {r["insight_id"]: r for f in V1.glob("*.judge.jsonl") for r in read_jsonl(f)}
    insights, cache = _load_insights(), {}
    picked = [k for k in key if marks[k["row"]][0] not in ("correct", "wrong")
              or marks[k["row"]][0] != k["judge_verdict"]]
    wb = Workbook()
    _guide(wb, "Final marks under the updated rules. The judge's verdict and reason are shown: "
               "it can be wrong too. Change your mark only if the rules say so.")
    ws = _sheet(wb, "Adjudicate", ["#", "Type", "Owner / value / deadline", "Claim",
                                   "What was really said", "Meeting summary", "Your first mark",
                                   "Judge said", "FINAL MARK", "If wrong: why", "Note"],
                [5, 12, 20, 40, 70, 45, 11, 40, 12, 14, 30])
    rows, adj_key = [], []
    for n, k in enumerate(picked, 1):
        j = v1[k["insight_id"]]
        first = marks[k["row"]][0] or "(blank)"
        rows.append([n, *_row_values(insights[k["insight_id"]], cache), first,
                     f"{j['verdict']} ({j['failure']}): {j['reason']}"])
        adj_key.append({"row": n, "insight_id": k["insight_id"], "condition": k["condition"]})
    _write_rows(ws, rows, mark_col=9)
    wb.save(ADJ)
    ADJ_KEY.write_text(json.dumps(adj_key, indent=1), encoding="utf-8")
    print(f"Wrote {ADJ} ({len(rows)} rows).")


def finalize() -> None:
    key = json.loads(AUDIT_KEY.read_text(encoding="utf-8"))
    first = read_marks(AUDIT, len(key), 6)
    adj_key = json.loads(ADJ_KEY.read_text(encoding="utf-8"))
    final = read_marks(ADJ, len(adj_key), 9)
    labels = {k["insight_id"]: {"mark": first[k["row"]][0], "why": first[k["row"]][1],
                                "condition": k["condition"]} for k in key}
    for k in adj_key:
        mark, why, note = final[k["row"]]
        if mark not in ("correct", "wrong", "unsure"):
            raise SystemExit(f"Adjudication row {k['row']} has no final mark.")
        labels[k["insight_id"]] |= {"mark": mark, "why": why, "note": note}
    DEV.write_text(json.dumps(labels, indent=1), encoding="utf-8")
    print(f"Wrote {DEV}: {sum(v['mark'] == 'wrong' for v in labels.values())} wrong, "
          f"{sum(v['mark'] == 'correct' for v in labels.values())} correct, "
          f"{sum(v['mark'] == 'unsure' for v in labels.values())} unsure.")


def build_test(seed: int = 11) -> None:
    if TEST.exists():
        raise SystemExit(f"{TEST} exists; delete it first for a fresh sample.")
    used = {k["insight_id"] for k in json.loads(AUDIT_KEY.read_text(encoding="utf-8"))}
    verdicts, insights = current_verdicts(), _load_insights()
    if not verdicts:
        raise SystemExit("No current judge labels; run python -m evals.judge first.")
    rng, picked = random.Random(seed), []
    for cond, n in TEST_PER_CONDITION.items():
        pool = [v for i, v in sorted(verdicts.items()) if v["condition"] == cond and i not in used]
        wrong = [v for v in pool if v["verdict"] == "wrong"]
        rest = [v for v in pool if v["verdict"] != "wrong"]
        k = min(len(wrong), n // 2)
        picked += rng.sample(wrong, k) + rng.sample(rest, n - k)
    rng.shuffle(picked)
    wb, cache = Workbook(), {}
    _guide(wb, "Blind check: compare each CLAIM with WHAT WAS REALLY SAID and the meeting summary.")
    ws = _sheet(wb, "Audit", ["#", "Type", "Owner / value / deadline", "Claim",
                              "What was really said", "Meeting summary", "YOUR MARK",
                              "If wrong: why", "Note"],
                [5, 12, 20, 45, 75, 45, 12, 14, 30])
    rows = [[n, *_row_values(insights[v["insight_id"]], cache)] for n, v in enumerate(picked, 1)]
    _write_rows(ws, rows, mark_col=7)
    wb.save(TEST)
    TEST_KEY.write_text(json.dumps([{"row": n, "insight_id": v["insight_id"],
                                     "condition": v["condition"]} for n, v in enumerate(picked, 1)],
                                   indent=1), encoding="utf-8")
    print(f"Wrote {TEST} ({len(rows)} rows).")


def score(which: str) -> None:
    if which == "dev":
        labels = json.loads(DEV.read_text(encoding="utf-8"))
    else:
        key = json.loads(TEST_KEY.read_text(encoding="utf-8"))
        marks = read_marks(TEST, len(key), 7)
        labels = {k["insight_id"]: {"mark": marks[k["row"]][0], "why": marks[k["row"]][1],
                                    "condition": k["condition"]} for k in key}
    verdicts = current_verdicts()
    pairs = [(i, v, verdicts[i]) for i, v in labels.items()
             if v["mark"] in ("correct", "wrong") and i in verdicts]
    agree = sum(v["mark"] == j["verdict"] for _, v, j in pairs)
    print(f"[{which}] judge agrees with human on {agree}/{len(pairs)} = {agree / len(pairs):.0%}")
    for say in ("wrong", "correct"):
        sub = [(v, j) for _, v, j in pairs if j["verdict"] == say]
        if sub:
            ok = sum(v["mark"] == say for v, _ in sub)
            print(f"  when judge says {say:7s}: human agrees {ok}/{len(sub)}")
    human_wrong = [(v, j) for _, v, j in pairs if v["mark"] == "wrong"]
    if human_wrong:
        caught = sum(j["verdict"] == "wrong" for _, j in human_wrong)
        print(f"  human-wrong insights the judge catches: {caught}/{len(human_wrong)}")


if __name__ == "__main__":
    p = argparse.ArgumentParser()
    g = p.add_mutually_exclusive_group(required=True)
    g.add_argument("--adjudicate", action="store_true")
    g.add_argument("--finalize", action="store_true")
    g.add_argument("--test", action="store_true")
    g.add_argument("--score", choices=["dev", "test"])
    a = p.parse_args()
    if a.adjudicate:
        adjudicate()
    elif a.finalize:
        finalize()
    elif a.test:
        build_test()
    else:
        score(a.score)

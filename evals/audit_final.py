"""Final blind audit of the judge, pre-committed as the reported agreement figure.

History (all three are reported): audit 1 found the rubric ambiguous (55%); the blind test of
the revised judge scored 52% (kappa 0.04) and the rater said he had marked too quickly. This
audit is blind and structured against rushing: for each row the rater answers the four rules
yes / no / n-a / unsure, and the verdict is derived (any "no" -> wrong). Whatever it shows is
the number the README reports.

  python -m evals.audit_final            # build data/labels/final_audit.xlsx (30 rows)
  python -m evals.audit_final --score
"""

import argparse
import json
import math
import random

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from evals.audit_sheet import AUDIT_KEY, TEST_KEY, _load_insights, _row_values, current_verdicts
from evals.judge import LABELS, RULES

SHEET, KEY = LABELS / "final_audit.xlsx", LABELS / "final_audit_key.json"
PER_CONDITION = {"farfield": 11, "headset": 7, "speaker_swap": 6, "clean": 6}
QUESTIONS = ["1. Said and still true?", "2. Right kind?", "3. Right person?", "4. Details match?"]
HELP = [
    "1. Said and still true: does what was really said support the claim, and does the summary "
    "NOT show it reversed or superseded later in the meeting? (A recap that was later changed = no.)",
    "2. Right kind: decision = explicitly agreed ('we decided', 'let's go with', clearly accepted). "
    "A 'maybe' answered only by 'yeah'/'mm' = no. Action item = someone committed. Open risk = "
    "still unresolved at the end of the meeting.",
    "3. Right person: action items only (otherwise n-a). Is the owner who actually took the task? "
    "'none named' is right only if nobody was given it.",
    "4. Details match: numbers, units, deadlines exactly as said (n-a if the claim has none).",
    "Use 'unsure' only if the excerpt and summary truly cannot settle it.",
    "Take about a minute per row: read the claim, then check each question against the excerpt.",
]
FIRST_COL = 7  # column G: first question


def build(seed: int = 23) -> None:
    if SHEET.exists():
        raise SystemExit(f"{SHEET} exists; it is the committed final audit.")
    used = {k["insight_id"] for path in (AUDIT_KEY, TEST_KEY)
            for k in json.loads(path.read_text(encoding="utf-8"))}
    verdicts, insights = current_verdicts(), _load_insights()
    rng, picked = random.Random(seed), []
    for cond, n in PER_CONDITION.items():
        pool = [v for i, v in sorted(verdicts.items()) if v["condition"] == cond and i not in used]
        wrong = [v for v in pool if v["verdict"] == "wrong"]
        rest = [v for v in pool if v["verdict"] != "wrong"]
        k = min(len(wrong), n // 2)
        picked += rng.sample(wrong, k) + rng.sample(rest, n - k)
    rng.shuffle(picked)

    wb = Workbook()
    guide = wb.active
    guide.title = "How to answer"
    guide.column_dimensions["A"].width = 120
    for i, line in enumerate(["Answer the four questions for every row (yes / no / n-a / unsure). "
                              "The verdict is worked out from your answers.", "", *HELP, "",
                              "Full rules:", *RULES.splitlines()], 1):
        guide.cell(i, 1, line).alignment = Alignment(wrap_text=True)
    guide.cell(1, 1).font = Font(bold=True)

    ws = wb.create_sheet("Audit")
    headers = ["#", "Type", "Owner / value / deadline", "Claim", "What was really said",
               "Meeting summary", *QUESTIONS, "Note"]
    widths = [5, 12, 20, 40, 70, 45, 12, 12, 12, 12, 30]
    for c, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(1, c, h)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDDDDD")
        cell.alignment = Alignment(wrap_text=True)
        ws.column_dimensions[cell.column_letter].width = w
    ws.freeze_panes = "A2"
    cache = {}
    for r, v in enumerate(picked, 2):
        for c, val in enumerate([r - 1, *_row_values(insights[v["insight_id"]], cache)], 1):
            ws.cell(r, c, val).alignment = Alignment(wrap_text=True, vertical="top")
        for c in range(FIRST_COL, FIRST_COL + 4):
            ws.cell(r, c).fill = PatternFill("solid", fgColor="FFF2CC")
    dv = DataValidation(type="list", formula1='"yes,no,n-a,unsure"', allow_blank=True)
    ws.add_data_validation(dv)
    first = ws.cell(1, FIRST_COL).column_letter
    last = ws.cell(1, FIRST_COL + 3).column_letter
    dv.add(f"{first}2:{last}{len(picked) + 1}")
    wb.save(SHEET)
    KEY.write_text(json.dumps([{"row": n, "insight_id": v["insight_id"], "condition": v["condition"]}
                               for n, v in enumerate(picked, 1)], indent=1), encoding="utf-8")
    print(f"Wrote {SHEET} ({len(picked)} rows).")


def derive(answers: list[str]) -> str:
    a = [(x or "").strip().lower() for x in answers]
    if "no" in a:
        return "wrong"
    if "unsure" in a or any(x not in ("yes", "n-a") for x in a[:2]):
        return "unsure"
    return "correct"


def wilson(k: int, n: int, z: float = 1.96) -> tuple[float, float]:
    p = k / n
    centre = (p + z * z / (2 * n)) / (1 + z * z / n)
    half = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n)) / (1 + z * z / n)
    return centre - half, centre + half


def score() -> None:
    key = json.loads(KEY.read_text(encoding="utf-8"))
    ws = load_workbook(SHEET, read_only=True)["Audit"]
    rows = {r[0]: r for r in ws.iter_rows(min_row=2, max_row=len(key) + 1,
                                          max_col=FIRST_COL + 4, values_only=True)}
    verdicts = current_verdicts()
    pairs = []
    for k in key:
        human = derive(list(rows[k["row"]][FIRST_COL - 1:FIRST_COL + 3]))
        if human != "unsure":
            pairs.append((human, verdicts[k["insight_id"]]["verdict"], k["condition"]))
    n = len(pairs)
    agree = sum(h == j for h, j, _ in pairs)
    lo, hi = wilson(agree, n)
    jw = sum(j == "wrong" for _, j, _ in pairs) / n
    hw = sum(h == "wrong" for h, _, _ in pairs) / n
    pe = jw * hw + (1 - jw) * (1 - hw)
    kappa = (agree / n - pe) / (1 - pe) if pe < 1 else float("nan")
    print(f"[final] judge agrees with human on {agree}/{n} = {agree / n:.0%} "
          f"(95% CI {lo:.0%}-{hi:.0%}), Cohen's kappa {kappa:.2f}; "
          f"{len(key) - n} unsure rows excluded")
    for say in ("wrong", "correct"):
        sub = [h for h, j, _ in pairs if j == say]
        if sub:
            print(f"  when judge says {say:7s}: human agrees {sum(h == say for h in sub)}/{len(sub)}")
    hum_wrong = [j for h, j, _ in pairs if h == "wrong"]
    if hum_wrong:
        print(f"  human-wrong insights the judge catches: {sum(j == 'wrong' for j in hum_wrong)}/{len(hum_wrong)}")


if __name__ == "__main__":
    p = argparse.ArgumentParser()
    p.add_argument("--score", action="store_true")
    score() if p.parse_args().score else build()

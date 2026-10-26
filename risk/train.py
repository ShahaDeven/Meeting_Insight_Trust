"""Train and evaluate the risk score that splits insights into auto-write vs review.

Protocol (fixed before reporting; revised 2026-10-25 after an external review):
- Leave-one-series-out CV (6 folds): every prediction is for a team the model never saw.
- PRIMARY model: logistic regression, fixed in advance (no selection on these CV results; it
  is also the simplest model). XGBoost is a secondary check only, never chosen by its score.
  Its sigmoid calibrator is fitted on series-grouped folds inside the training series, so no
  calibration or model-selection decision ever sees the held-out series.
- Uncertainty: bootstrap over whole recording series (meetings in a series are not independent).
- Per-series results are reported for all six held-out series, not only pooled.
Baselines: the extractor's own confidence, and recognizer confidence on the quote.

Main metric: share of wrong insights caught when the riskiest X% go to human review, and the
error rate left among auto-written insights. Labels are judge v2 verdicts (a strict reviewer,
66% agreement with a blind human; see CLAUDE.md). The score is also checked against the
human-labelled insights (adjudicated dev set + final audit).

  python -m risk.train
Out: data/results/risk_scores.csv, risk_eval.csv, risk_calibration.csv, report/risk_*.png
"""

import json
import random

import numpy as np
import pandas as pd
from sklearn.calibration import CalibratedClassifierCV
from sklearn.model_selection import GroupKFold
from sklearn.linear_model import LogisticRegression
from sklearn.metrics import roc_auc_score
from sklearn.pipeline import make_pipeline
from sklearn.preprocessing import StandardScaler
from xgboost import XGBClassifier

from core.config import REPORT, RESULTS
from evals.audit_final import KEY as FINAL_KEY, SHEET as FINAL_SHEET, FIRST_COL, derive
from evals.audit_sheet import DEV

BUDGETS = (0.10, 0.20, 0.30)


def feature_columns(df: pd.DataFrame) -> list[str]:
    skip = {"insight_id", "meeting", "series", "condition", "run", "wrong", "failure"}
    return [c for c in df.columns if c not in skip]


PRIMARY = "logistic"
SCORES = ["base_model_conf", "base_quote_conf", "logistic", "xgboost"]


def models(groups: pd.Series | None = None, X: pd.DataFrame | None = None,
           y: pd.Series | None = None) -> dict:
    """Fresh models. Pass the training rows' series to get series-grouped calibration folds."""
    cv = (list(GroupKFold(n_splits=3).split(X, y, groups=groups)) if groups is not None else 3)
    return {
        "logistic": make_pipeline(StandardScaler(), LogisticRegression(max_iter=2000, C=0.5)),
        "xgboost": CalibratedClassifierCV(
            XGBClassifier(n_estimators=200, max_depth=3, learning_rate=0.05, subsample=0.8,
                          colsample_bytree=0.8, eval_metric="logloss", verbosity=0),
            method="sigmoid", cv=cv),
    }


def cross_val(df: pd.DataFrame, cols: list[str]) -> pd.DataFrame:
    out = df[["insight_id", "meeting", "series", "condition", "wrong"]].copy()
    out["base_model_conf"] = -df.model_confidence     # higher score = riskier
    out["base_quote_conf"] = -df.quote_conf_min
    for name in ("logistic", "xgboost"):
        out[name] = np.nan
    for s in sorted(df.series.unique()):
        train, test = df.series != s, df.series == s
        X, y = df.loc[train, cols].reset_index(drop=True), df.loc[train, "wrong"].reset_index(drop=True)
        for name, m in models(df.loc[train, "series"].reset_index(drop=True), X, y).items():
            m.fit(X, y)
            out.loc[test, name] = m.predict_proba(df.loc[test, cols])[:, 1]
    return out


def caught(y: np.ndarray, score: np.ndarray, budget: float) -> tuple[float, float]:
    """(share of wrong caught in the riskiest `budget`, error rate among the rest)."""
    order = np.argsort(-score, kind="stable")
    k = int(round(budget * len(y)))
    review, auto = y[order[:k]], y[order[k:]]
    return review.sum() / max(y.sum(), 1), auto.mean() if len(auto) else 0.0


def evaluate(scores: pd.DataFrame, n_boot: int = 2000, seed: int = 0) -> pd.DataFrame:
    """Pooled metrics; 95% CIs from resampling whole series (6 groups, so they are wide)."""
    series = sorted(scores.series.unique())
    by_series = {s: g for s, g in scores.groupby("series")}
    rows = []
    for col in SCORES:
        for subset, d in [("all", scores), ("asr_only", scores[scores.condition.isin(["headset", "farfield"])])]:
            y, s = d.wrong.to_numpy(), d[col].to_numpy()
            row = {"score": col, "subset": subset, "n": len(d), "auroc": roc_auc_score(y, s)}
            for b in BUDGETS:
                row[f"caught@{int(b * 100)}"], row[f"auto_err@{int(b * 100)}"] = caught(y, s, b)
            if subset == "all":
                rng = random.Random(seed)  # same resamples for every score, so CIs are comparable
                aucs, c20 = [], []
                for _ in range(n_boot):
                    g = pd.concat([by_series[rng.choice(series)] for _ in series])
                    if g.wrong.nunique() == 2:
                        aucs.append(roc_auc_score(g.wrong, g[col]))
                        c20.append(caught(g.wrong.to_numpy(), g[col].to_numpy(), 0.2)[0])
                aucs.sort()
                c20.sort()
                q = lambda v, p: v[int(p * len(v))]  # noqa: E731
                row |= {"auroc_lo": q(aucs, 0.025), "auroc_hi": q(aucs, 0.975),
                        "caught@20_lo": q(c20, 0.025), "caught@20_hi": q(c20, 0.975)}
            rows.append(row)
    return pd.DataFrame(rows)


def paired_gain(scores: pd.DataFrame, a: str = PRIMARY, b: str = "base_model_conf",
                n_boot: int = 2000, seed: int = 0) -> dict:
    """Score a minus score b on the same series resamples (AUROC and caught at 20%)."""
    series = sorted(scores.series.unique())
    by_series = {s: g for s, g in scores.groupby("series")}
    rng = random.Random(seed)
    d_auc, d_c20 = [], []
    for _ in range(n_boot):
        g = pd.concat([by_series[rng.choice(series)] for _ in series])
        if g.wrong.nunique() == 2:
            y = g.wrong.to_numpy()
            d_auc.append(roc_auc_score(y, g[a]) - roc_auc_score(y, g[b]))
            d_c20.append(caught(y, g[a].to_numpy(), 0.2)[0] - caught(y, g[b].to_numpy(), 0.2)[0])
    y = scores.wrong.to_numpy()
    out = {"a": a, "b": b,
           "auroc_gain": roc_auc_score(y, scores[a]) - roc_auc_score(y, scores[b]),
           "caught20_gain": caught(y, scores[a].to_numpy(), 0.2)[0] - caught(y, scores[b].to_numpy(), 0.2)[0]}
    for name, v in (("auroc_gain", sorted(d_auc)), ("caught20_gain", sorted(d_c20))):
        out[f"{name}_lo"], out[f"{name}_hi"] = v[int(0.025 * len(v))], v[int(0.975 * len(v))]
    return out


def per_series(scores: pd.DataFrame) -> pd.DataFrame:
    """Each held-out series on its own: AUROC and wrong insights caught at a 20% review budget."""
    rows = []
    for s, d in scores.groupby("series"):
        row = {"series": s, "insights": len(d), "wrong": int(d.wrong.sum())}
        for col in SCORES:
            row[f"auroc_{col}"] = roc_auc_score(d.wrong, d[col]) if d.wrong.nunique() == 2 else np.nan
            row[f"caught20_{col}"] = caught(d.wrong.to_numpy(), d[col].to_numpy(), 0.2)[0]
        rows.append(row)
    return pd.DataFrame(rows)


def calibration(scores: pd.DataFrame, col: str, bins: int = 8) -> pd.DataFrame:
    d = scores[[col, "wrong"]].copy()
    d["bin"] = pd.qcut(d[col], q=bins, duplicates="drop")
    return d.groupby("bin", observed=True).agg(predicted=(col, "mean"), observed=("wrong", "mean"),
                                               n=("wrong", "size")).reset_index(drop=True)


def human_check(scores: pd.DataFrame) -> pd.DataFrame:
    """AUROC of each score against the human labels (adjudicated dev set + final audit)."""
    human = {i: v["mark"] for i, v in json.loads(DEV.read_text(encoding="utf-8")).items()}
    from openpyxl import load_workbook
    key = json.loads(FINAL_KEY.read_text(encoding="utf-8"))
    ws = load_workbook(FINAL_SHEET, read_only=True)["Audit"]
    rows = {r[0]: r for r in ws.iter_rows(min_row=2, max_row=len(key) + 1,
                                          max_col=FIRST_COL + 4, values_only=True)}
    for k in key:
        human[k["insight_id"]] = derive(list(rows[k["row"]][FIRST_COL - 1:FIRST_COL + 3]))
    # Speaker-swap insights were re-extracted after the dose fix; their ids now name different
    # insights from the ones the humans labelled, so they are left out of this check.
    human = {i: m for i, m in human.items() if ".speaker_swap." not in i}
    d = scores[scores.insight_id.isin(human)].copy()
    d["human_wrong"] = d.insight_id.map(lambda i: int(human[i] == "wrong"))
    d = d[d.insight_id.map(lambda i: human[i] in ("correct", "wrong"))]
    return pd.DataFrame([{"score": c, "n": len(d), "human_wrong": int(d.human_wrong.sum()),
                          "auroc_vs_human": roc_auc_score(d.human_wrong, d[c])}
                         for c in ["base_model_conf", "base_quote_conf", "logistic", "xgboost"]])


def plots(scores: pd.DataFrame, cal: pd.DataFrame, best: str) -> None:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    REPORT.mkdir(parents=True, exist_ok=True)
    y = scores.wrong.to_numpy()
    xs = np.linspace(0, 1, 101)
    fig, ax = plt.subplots(figsize=(6, 4))
    for col, label in [(best, f"risk model ({best})"), ("base_model_conf", "extractor confidence"),
                       ("base_quote_conf", "ASR confidence on quote")]:
        ax.plot(xs * 100, [caught(y, scores[col].to_numpy(), b)[0] * 100 for b in xs], label=label)
    ax.plot([0, 100], [0, 100], ":", color="grey", label="random")
    ax.set(xlabel="% of insights sent to human review (riskiest first)",
           ylabel="% of wrong insights caught", title="Review budget vs wrong writes caught")
    ax.legend()
    fig.tight_layout()
    fig.savefig(REPORT / "risk_catch_curve.png", dpi=150)
    fig, ax = plt.subplots(figsize=(4.5, 4.5))
    ax.plot([0, 1], [0, 1], ":", color="grey")
    ax.plot(cal.predicted, cal.observed, "o-")
    ax.set(xlabel="predicted risk", ylabel="observed wrong rate", title=f"Calibration ({best})")
    fig.tight_layout()
    fig.savefig(REPORT / "risk_calibration.png", dpi=150)


def main() -> None:
    df = pd.read_csv(RESULTS / "risk_features.csv")
    cols = feature_columns(df)
    scores = cross_val(df, cols)
    scores.to_csv(RESULTS / "risk_scores.csv", index=False)
    ev = evaluate(scores)
    ev.to_csv(RESULTS / "risk_eval.csv", index=False)
    best = PRIMARY  # fixed in advance, never chosen from these results
    ps = per_series(scores)
    ps.to_csv(RESULTS / "risk_per_series.csv", index=False)
    gain = paired_gain(scores)
    (RESULTS / "risk_gain.json").write_text(json.dumps(gain, indent=1), encoding="utf-8")
    cal = calibration(scores, best)
    cal.to_csv(RESULTS / "risk_calibration.csv", index=False)
    hc = human_check(scores)
    hc.to_csv(RESULTS / "risk_human_check.csv", index=False)
    plots(scores, cal, best)

    pd.set_option("display.width", 200)
    fmt = {c: "{:.0%}".format for c in ev.columns if c.startswith(("caught", "auto_err"))}
    print(f"{len(df)} insights, {df.wrong.mean():.1%} wrong; {len(cols)} features; "
          f"leave-one-series-out CV.\n")
    print(ev.to_string(index=False, formatters=fmt | {"auroc": "{:.3f}".format,
                                                      "auroc_lo": "{:.3f}".format, "auroc_hi": "{:.3f}".format}))
    print(f"\nPaired gain over extractor confidence: AUROC {gain['auroc_gain']:+.3f} "
          f"[{gain['auroc_gain_lo']:+.3f}, {gain['auroc_gain_hi']:+.3f}], caught@20 "
          f"{gain['caught20_gain']:+.1%} [{gain['caught20_gain_lo']:+.1%}, {gain['caught20_gain_hi']:+.1%}]")
    print(f"\nPer held-out series (primary = {PRIMARY}):")
    print(ps.to_string(index=False, float_format=lambda v: f"{v:.2f}"))
    print(f"\nCalibration ({best}):\n{cal.round(3).to_string(index=False)}")
    print(f"\nAgainst human labels:\n{hc.round(3).to_string(index=False)}")
    lr = models()["logistic"].fit(df[cols], df.wrong)
    coef = pd.Series(lr[-1].coef_[0], index=cols).sort_values()
    print(f"\nLogistic coefficients (standardized; + = riskier):\n{coef.round(2).to_string()}")


if __name__ == "__main__":
    main()
